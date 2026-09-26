"""Tirage de l'échantillon de contrôle manuel du corpus filtré (ticket #27, tâche 7).

Produit une grille Excel à remplir à la main :
  A  50 passages conservés, tirés au hasard : le passage est-il utile et compréhensible ?
  B  20 passages retirés, répartis entre R1, R2 et R3 : le retrait était-il justifié ?

La colonne à encoder n'étant pas encore choisie par le domaine D2, chaque passage est
contrôlé sous ses deux formes, text et chunk_text, avec un verdict pour chacune.

Le tirage utilise une graine fixe : il est identique à chaque exécution.
Prérequis : avoir lancé scripts/prepare_corpus.py.

Usage, depuis la racine du dépôt :
    python scripts/echantillon_controle.py            # refuse d'écraser une grille existante
    python scripts/echantillon_controle.py --ecraser  # régénère la grille, annotations perdues
"""
import argparse
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

RACINE = Path(__file__).resolve().parent.parent
FILTRE = RACINE / "data/processed/travail_emploi_filtre.parquet"
RETIRES = RACINE / "data/processed/passages_retires.csv"
GRILLE = RACINE / "docs/corpus/controle_echantillon.xlsx"

GRAINE = 27          # numéro du ticket
TAILLE_A = 50
TAILLE_B = {"R1": 7, "R2": 7, "R3": 6}

VERDICTS = ["OK", "problème"]
TYPES_TEXT = [
    "coupé en milieu de phrase",
    "fragment peu compréhensible",
    "bruit de navigation ou liens",
    "mauvaise section (context)",
    "tableau aplati",
    "hors sujet",
    "retrait injustifié",
    "autre",
]
# chunk_text contient text : il hérite de ses problèmes, et en ajoute qui lui sont propres
TYPES_CHUNK_TEXT = [
    "introduction vide",
    "introduction sans rapport avec le passage",
    "titre sans rapport avec le passage",
    "trop long pour l'encodeur",
    *[t for t in TYPES_TEXT if t != "autre"],
    "autre",
]
LISTES = {
    "verdict_text": VERDICTS, "type_probleme_text": TYPES_TEXT,
    "verdict_chunk_text": VERDICTS, "type_probleme_chunk_text": TYPES_CHUNK_TEXT,
}
A_REMPLIR = [*LISTES, "commentaire"]

ACCENT = "1F4E5F"    # couleur d'accent de la charte de l'équipe (docs/latex/preambule.tex)


def coupure(texte):
    """Indique si le passage commence en milieu de phrase (minuscule) et s'il finit sans ponctuation forte."""
    t = texte.strip()
    debut = bool(pd.Series([t]).str.match(r"^[a-zàâçéèêëîïôûùüÿœ]").iloc[0])
    fin = not bool(pd.Series([t]).str.contains(r"[.!?»:;)]$").iloc[0])
    return {(True, True): "début et fin", (True, False): "début",
            (False, True): "fin", (False, False): "aucune"}[(debut, fin)]


def tirer():
    garde = pd.read_parquet(FILTRE)
    retires = pd.read_csv(RETIRES, sep=";", encoding="utf-8-sig")
    garde["section"] = garde["context"].apply(lambda x: x[0])

    a = garde.sample(n=TAILLE_A, random_state=GRAINE)

    b = pd.concat([retires[retires["regle"] == r].sample(n=n, random_state=GRAINE)
                   for r, n in TAILLE_B.items()])
    # le CSV des retirés ne contient pas chunk_text : on le reprend du fichier brut
    brut = pd.read_parquet(RACINE / "data/raw/travail-emploi-20260911/travail_emploi_part_0.parquet",
                           columns=["chunk_id", "chunk_text"])
    b = b.merge(brut, on="chunk_id", how="left")

    def mettre_en_forme(ech, colonnes_propres):
        ech = ech.copy()
        ech.insert(0, "n", range(1, len(ech) + 1))
        ech["longueur_text"] = ech["text"].str.strip().str.len()
        ech["coupure"] = ech["text"].apply(coupure)
        ech["longueur_chunk_text"] = ech["chunk_text"].str.strip().str.len()
        for col in A_REMPLIR:
            ech[col] = ""
        return ech[["n", "chunk_id", "title", "section", "chunk_index", *colonnes_propres,
                    "longueur_text", "coupure", "text", "verdict_text", "type_probleme_text",
                    "longueur_chunk_text", "chunk_text", "verdict_chunk_text", "type_probleme_chunk_text",
                    "commentaire"]]

    return mettre_en_forme(a, ["doublon_text"]), mettre_en_forme(b, ["regle", "motif"])


def feuille(classeur, titre, df, largeurs):
    ws = classeur.create_sheet(titre)
    ws.append(list(df.columns))
    for ligne in df.itertuples(index=False):
        ws.append(list(ligne))

    entete, a_remplir = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="FFF7E0")
    for cellule in ws[1]:
        cellule.font = entete
        cellule.fill = PatternFill("solid", fgColor=ACCENT)
        cellule.alignment = Alignment(vertical="center", wrap_text=True)
    for i, col in enumerate(df.columns, start=1):
        lettre = ws.cell(row=1, column=i).column_letter
        ws.column_dimensions[lettre].width = largeurs.get(col, 14)
        for cellule in ws[lettre][1:]:
            cellule.alignment = Alignment(vertical="top", wrap_text=True)
            if col in A_REMPLIR:
                cellule.fill = a_remplir
    ws.freeze_panes = "B2"

    # listes déroulantes sur les colonnes à remplir, alimentées par l'onglet Listes :
    # Excel limite à 255 caractères une liste écrite directement dans la validation
    for i, (col, valeurs) in enumerate(LISTES.items(), start=1):
        lettre = ws.cell(row=1, column=df.columns.get_loc(col) + 1).column_letter
        source = ws.cell(row=1, column=i).column_letter
        dv = DataValidation(type="list", formula1=f"Listes!${source}$2:${source}${len(valeurs) + 1}",
                            allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f"{lettre}2:{lettre}{len(df) + 1}")


def listes(classeur):
    ws = classeur.create_sheet("Listes")
    for i, (col, valeurs) in enumerate(LISTES.items(), start=1):
        ws.cell(row=1, column=i, value=col).font = Font(bold=True)
        for j, v in enumerate(valeurs, start=2):
            ws.cell(row=j, column=i, value=v)
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = 40


def consignes(classeur):
    ws = classeur.active
    ws.title = "Consignes"
    lignes = [
        "Contrôle manuel du corpus filtré — ticket #27, tâche 7",
        "",
        f"Échantillon tiré par scripts/echantillon_controle.py, graine {GRAINE}, à partir du corpus produit par scripts/prepare_corpus.py.",
        "",
        f"Feuille A_gardes : {TAILLE_A} passages conservés, tirés au hasard.",
        "   Question : le passage est-il un contenu utile, compréhensible et cohérent avec sa fiche et sa section ?",
        f"Feuille B_retires : {sum(TAILLE_B.values())} passages retirés ({', '.join(f'{n} {r}' for r, n in TAILLE_B.items())}).",
        "   Question : le retrait était-il justifié ?",
        "",
        "Chaque passage est présenté sous ses deux formes, car le domaine D2 n'a pas encore choisi la colonne à encoder :",
        "   text : le contenu du passage seul",
        "   chunk_text : titre de la fiche + introduction + text",
        "Donner un verdict pour chacune des deux formes.",
        "",
        "Colonne coupure (calculée automatiquement) : le passage commence-t-il en milieu de phrase (minuscule initiale)",
        "et finit-il sans ponctuation forte ? Valeurs : début, fin, début et fin, aucune. La coupure est un défaut du",
        "découpage de l'éditeur, commun à text et chunk_text ; le verdict porte sur l'utilité du passage malgré elle.",
        "",
        "Colonnes à remplir (fond jaune) :",
        "   verdict_text, verdict_chunk_text : OK ou problème (liste déroulante)",
        "   type_probleme_text, type_probleme_chunk_text : si problème, choisir dans la liste déroulante",
        "   commentaire : précision libre",
        "",
        "Types de problème pour text :",
        *[f"   {t}" for t in TYPES_TEXT],
        "",
        "Types de problème pour chunk_text :",
        *[f"   {t}" for t in TYPES_CHUNK_TEXT],
        "   (« trop long pour l'encodeur » : multilingual-e5-small tronque au-delà de 512 tokens ; le nombre",
        "    exact de tokens figure dans les colonnes tokens_text et tokens_chunk_text du corpus filtré)",
    ]
    for texte in lignes:
        ws.append([texte])
    ws["A1"].font = Font(bold=True, size=13, color=ACCENT)
    ws.column_dimensions["A"].width = 120


def main():
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--ecraser", action="store_true", help="régénérer la grille même si elle existe déjà")
    args = parser.parse_args()

    if GRILLE.exists() and not args.ecraser:
        raise SystemExit(f"{GRILLE.relative_to(RACINE)} existe déjà : relancer avec --ecraser pour la régénérer "
                         "(les annotations seraient perdues)")

    a, b = tirer()
    classeur = Workbook()
    consignes(classeur)
    largeurs = {"n": 5, "chunk_id": 20, "title": 24, "section": 24, "chunk_index": 8, "doublon_text": 9,
                "regle": 7, "motif": 24, "longueur_text": 9, "coupure": 12, "text": 60,
                "verdict_text": 11, "type_probleme_text": 22,
                "longueur_chunk_text": 9, "chunk_text": 70,
                "verdict_chunk_text": 11, "type_probleme_chunk_text": 22,
                "commentaire": 30}
    feuille(classeur, "A_gardes", a, largeurs)
    feuille(classeur, "B_retires", b, largeurs)
    listes(classeur)
    GRILLE.parent.mkdir(parents=True, exist_ok=True)
    classeur.save(GRILLE)
    print(f"grille écrite : {GRILLE.relative_to(RACINE)} ({len(a)} passages conservés, {len(b)} retirés)")


if __name__ == "__main__":
    main()
